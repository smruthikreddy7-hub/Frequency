import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def create_pitch_document():
    doc = docx.Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Style definitions
    # Primary Palette
    DARK_BLUE = RGBColor(15, 23, 42)     # #0f172a
    EMERALD = RGBColor(16, 185, 129)     # #10b981
    SLATE_GRAY = RGBColor(71, 85, 105)   # #475569
    TEXT_DARK = RGBColor(30, 41, 59)     # #1e293b

    # =========================================================================
    # TITLE & HEADER
    # =========================================================================
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_p.add_run("🎙️ FREQUENCY")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = DARK_BLUE

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Cross-Sense Multimodal AI Companion for Human Energy & Pattern Discovery")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(14)
    sub_run.font.bold = True
    sub_run.font.color.rgb = EMERALD

    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run("Complete Jury Pitch Deck, Technical Architecture & Live Demo Guide\n" + "—" * 65)
    meta_run.font.name = "Calibri"
    meta_run.font.size = Pt(10)
    meta_run.font.color.rgb = SLATE_GRAY

    # =========================================================================
    # SECTION 1: THE ELEVATOR PITCH
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("1. The 30-Second Elevator Pitch (The Hook)")
    r1.font.color.rgb = DARK_BLUE

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(
        "\"Most AI assistants only listen to what you type, while fitness wearables only track raw biometric numbers without understanding your thoughts. "
        "FREQUENCY is a privacy-first Cross-Sense AI that bridges this gap.\n\n"
        "It analyzes what you say (linguistic semantics), how you sound (vocal acoustics), and how you live (sleep & workload context) "
        "against your personal historical baseline to uncover your true focus, fatigue, and energy patterns—running 100% privately on your local machine with zero cloud data leakage.\""
    )
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.italic = True

    # =========================================================================
    # SECTION 2: THE PROBLEM STATEMENT
    # =========================================================================
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("2. The Problem We Are Solving")
    r2.font.color.rgb = DARK_BLUE

    problems = [
        ("The 'Masked Strain' Dilemma (Polite Words vs. Real Fatigue)",
         "When individuals experience burnout or cognitive overload, they frequently mask their distress with polite phrases (e.g., 'I am fine, just finishing up work'). Traditional chatbots take text at face value and miss the underlying crisis."),
        ("The 'Disconnected Data' Silo",
         "Wearables collect sleep and heart rate data in silos without knowing what you are working on or how you feel. Journaling apps collect text reflections without acoustic or physiological context. No tool connects how you sleep to how you speak and think."),
        ("The Critical Privacy Barrier",
         "Mental reflections, voice notes, and emotional vulnerabilities are the most sensitive personal data. Uploading raw voice recordings and journals to centralized third-party cloud servers poses severe data leak liabilities.")
    ]

    for title, desc in problems:
        p = doc.add_paragraph()
        r_title = p.add_run(f"• {title}: ")
        r_title.font.bold = True
        r_title.font.color.rgb = DARK_BLUE
        r_desc = p.add_run(desc)
        r_desc.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SECTION 3: THE 6-LAYER ARCHITECTURE
    # =========================================================================
    h3 = doc.add_heading(level=1)
    r3 = h3.add_run("3. The 6-Layer Multi-Modal Architecture")
    r3.font.color.rgb = DARK_BLUE

    layers = [
        ("Layer 1: Text / Linguistic Modality", "Custom local NLP engine extracting sentiment valence (-1.0 to +1.0), emotional affect categories (Exhaustion, Stress, Joy, Calm), thematic keywords (#workload, #sleep), and cognitive load markers."),
        ("Layer 2: Voice Acoustic Modality", "Local psychoacoustic analyzer calculating fundamental pitch F0, pitch inflection variance (Hz), vocal energy (RMS amplitude), speaking rate (WPM), and silence/pause duration ratios via autocorrelation and frame energy thresholding."),
        ("Layer 3: Daily Context Modality", "Lifestyle check-in engine tracking sleep duration & quality, workload intensity (Light, Normal, Heavy, Burnout Risk), physical activity, and self-reported mood scores (1-10 scale)."),
        ("Layer 4: Personal Baseline Engine", "7-day rolling statistical baseline engine computing historical personal distributions and calculating exact session percentage deviations (Δ%) rather than generic population averages."),
        ("Layer 5: Cross-Sense Fusion Matrix", "Synthesizes multi-modal streams into a Composite Energy Index (0-100%), Cognitive Focus Index (0-100%), and Sense Alignment Score (0-100%), classifying holistic Pattern Archetypes (e.g., Sleep-Deprived Sluggishness, Restored High-Flow Harmony)."),
        ("Layer 6: Conversational Reasoning Engine", "Local LLM reasoning bridge (Ollama / Qwen3 / Llama 3) compiling structured evidence packages into empathetic, non-diagnostic companion dialogue with zero cloud exposure.")
    ]

    for title, desc in layers:
        p = doc.add_paragraph()
        r_title = p.add_run(f"[{title}]: ")
        r_title.font.bold = True
        r_title.font.color.rgb = EMERALD
        r_desc = p.add_run(desc)
        r_desc.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SECTION 4: COMPLETE TECH STACK TABLE
    # =========================================================================
    h4 = doc.add_heading(level=1)
    r4 = h4.add_run("4. Complete Technology Stack")
    r4.font.color.rgb = DARK_BLUE

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Component Layer", "Technologies / Libraries", "Architectural Rationale"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "1e293b")
        p = hdr_cells[i].paragraphs[0]
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

    stack_rows = [
        ("Backend Framework", "Python 3.10, Flask", "Lightweight application factory, fast REST API, low-latency Server-Sent Events (SSE) token streaming."),
        ("Local AI Reasoning", "Ollama (Qwen3:8B / Llama3:8B)", "Zero cloud dependencies, 100% private on-device inference, 4096-token context window, streaming token pipeline."),
        ("Audio & Acoustics", "soundfile, scipy.signal, numpy", "Deterministic autocorrelation F0 pitch estimation, RMS amplitude extraction, pause ratio analysis with 0 cloud latency."),
        ("Speech-to-Text (STT)", "SpeechRecognition + Web Audio API", "In-browser 16kHz 16-bit PCM WAV encoding with dual client/server transcription fallbacks."),
        ("Local Database", "SQLite3 (frequency.db)", "Zero-configuration embedded relational store for daily context, multi-modal session logs, and baseline distributions."),
        ("Frontend Web UI", "Vanilla ES6+ JS, HTML5, Modern CSS3", "Fast zero-build frontend, responsive dark-mode design system, interactive Signal Telemetry drawer, 7-Day Trends Dashboard.")
    ]

    for comp, tech, rat in stack_rows:
        row_cells = table.add_row().cells
        row_cells[0].text = comp
        row_cells[1].text = tech
        row_cells[2].text = rat
        for cell in row_cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SECTION 5: JURY STRATEGY & Q&A GUIDE
    # =========================================================================
    h5 = doc.add_heading(level=1)
    r5 = h5.add_run("5. Jury Defense & Architecture Strategy")
    r5.font.color.rgb = DARK_BLUE

    qas = [
        ("Q: 'Is there a database in this application?'",
         "A: YES! We use an active, embedded SQLite relational database (frequency.db). It stores daily lifestyle context (sleep, workload, mood) and historical session telemetry. This persistence layer powers the 7-Day Rolling Baseline Engine and the Trends Dashboard."),
        ("Q: 'Why is there no traditional cloud login/signup portal? Is it mandatory?'",
         "A: NO, it is intentionally not included because FREQUENCY is built on a Local-First, Zero-Trust Privacy Philosophy (the same paradigm used by Apple Health, Obsidian, and Signal):\n"
         "  1. Zero-Trust Biometric Privacy: Voice pitch, emotional reflections, and sleep deficits are sensitive psychiatric/biometric data. Centralizing them behind cloud passwords creates severe hacking and leak liabilities.\n"
         "  2. Complete Psychological Safety: Users are far more honest when they know audio never leaves their hardware.\n"
         "  3. Zero Infrastructure Cost: $0 server hosting fees; the client hardware executes the computation.\n"
         "  4. Enterprise Ready: The data layer is decoupled (schema.py & database.py include user_id), making it trivial to attach OAuth2/JWT if deploying for clinical or enterprise multi-user environments.")
    ]

    for q, a in qas:
        p = doc.add_paragraph()
        rq = p.add_run(f"{q}\n")
        rq.font.bold = True
        rq.font.color.rgb = DARK_BLUE
        ra = p.add_run(a)
        ra.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SECTION 6: LIVE DEMO SCRIPT
    # =========================================================================
    h6 = doc.add_heading(level=1)
    r6 = h6.add_run("6. Step-by-Step 3-Minute Live Demo Script")
    r6.font.color.rgb = DARK_BLUE

    steps = [
        ("Step 1: Set Daily Context & Baseline (30s)", "Click '🌙 Context' in the header. Set Sleep to 4.5 hours and Workload to 'Heavy'. Click 'Update Context'. Open the Evidence Drawer to demonstrate how the Baseline Engine immediately flags a -39% sleep deficit."),
        ("Step 2: Record a Spoken Voice Note (1 min)", "Click the 'Microphone (🎙️)' button and speak: 'I’ve been grinding on this code refactor since 8 AM, but my tests keep failing and I can’t focus.' Stop recording and show the live transcript and the Composite Energy Index (e.g. 48%)."),
        ("Step 3: Dual-Sense AI Reasoning (1 min)", "Point out how the AI assistant responds directly to the code refactoring problem first, weaves in how the 4.5h sleep deficit and subdued vocal energy explain the mental friction, and offers a practical micro-restorative step."),
        ("Step 4: Expand Telemetry & Trends (30s)", "Click '📊 Signal Telemetry' under the message to show the exact mathematical readings (Pitch Hz, RMS energy, speech rate, baseline shift). Click '📈 Trends' in the header to show the 7-day multi-modal stability chart.")
    ]

    for step, desc in steps:
        p = doc.add_paragraph()
        r_step = p.add_run(f"{step}: ")
        r_step.font.bold = True
        r_step.font.color.rgb = EMERALD
        r_desc = p.add_run(desc)
        r_desc.font.color.rgb = TEXT_DARK

    # Save Document
    doc_path = os.path.join(os.path.dirname(__file__), "FREQUENCY_Pitch_Deck_and_Technical_Guide.docx")
    doc.save(doc_path)
    print(f"Successfully generated Word document at: {doc_path}")
    return doc_path

if __name__ == "__main__":
    create_pitch_document()
